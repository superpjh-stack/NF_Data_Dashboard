"""Read-only source extraction with a reproducible, atomically replaced SQLite snapshot."""
from __future__ import annotations
import argparse
from collections import Counter
from concurrent.futures import ProcessPoolExecutor
from datetime import date, datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import sqlite3
import unicodedata
import warnings

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / 'data/needsfood.sqlite'
PARSER_VERSION = '1.1'

def clean(value):
    if value is None: return ''
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFC', str(value))).strip()

def number(value):
    if isinstance(value, bool): return None
    if isinstance(value, (int, float)) and math.isfinite(value): return float(value)
    # Text is preserved; never reinterpret strings, limits or missing cells as numbers.
    return None

def source_root():
    for project in Path('/Users/gerardo92/Desktop').iterdir():
        if project.name.startswith('06 '):
            for p in project.iterdir():
                if clean(p.name) == '니즈푸드 데이터': return p
    raise FileNotFoundError('니즈푸드 데이터 폴더를 찾을 수 없습니다. --source로 지정하세요.')

def category(path):
    s = clean(str(path))
    if '_MES_마스터후보' in s: return '마스터 후보'
    if '주간식단표' in s: return '주간식단표'
    if '생산일지,원료수불부' in s:
        if '원료수불부' in clean(path.name): return '원료수불부'
        if '생산일지' in clean(path.name): return '생산일지'
        return '생산·수불 기타'
    if '/레시피/' in s: return '조리계획서'
    if '열량및알러지' in s: return '영양·알러지'
    if '식단표' in s: return '식단표'
    return '기타'

def json_value(v):
    if isinstance(v, (date, datetime)): return v.isoformat()
    if isinstance(v, float) and not math.isfinite(v): return str(v)
    return v

def workbook_sheets(path):
    """Yield (sheet name, rows with physical 1-based row indices)."""
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            w = load_workbook(path, read_only=True, data_only=True)
        try:
            for s in w:
                # Ignore inflated formatting dimensions; reset and stream actual XML rows.
                s.reset_dimensions()
                yield s.title, ((i, list(row)) for i, row in enumerate(s.iter_rows(values_only=True), 1))
        finally: w.close()
    else:
        from pyxlsb import open_workbook
        with open_workbook(str(path)) as w:
            for title in w.sheets:
                with w.get_sheet(title) as s:
                    def rows(sheet=s):
                        for cells in sheet.rows(sparse=True):
                            if not cells: continue
                            row = [None] * (max(c.c for c in cells) + 1)
                            for c in cells: row[c.c] = c.v
                            yield cells[0].r + 1, row
                    yield title, rows()

def date_in_row(row):
    for v in row:
        if isinstance(v, datetime): return v.date().isoformat()
        if isinstance(v, date): return v.isoformat()
    return None

def parse_sheet(rows, kind, sheet):
    records, preview, headers = [], [], None
    work_date, meal, menu = None, '', ''
    recognized = 0
    for ri, row in rows:
        if ri <= 12:
            preview.append({'row':ri, 'cells':[json_value(v) for v in row[:24]]})
        labels = [clean(v) for v in row]
        compact = [re.sub(r'\s+', '', s) for s in labels]
        # Production date is an actual typed workbook date, not a folder guess.
        if any('생산일자' in x or x == '일자' for x in compact):
            work_date = date_in_row(row) or work_date
            headers = None; menu = ''; continue
        meal_marker = False
        for v in labels[:2]:
            if v in ('조식', '중식', '석식', '간식'):
                meal = v; menu = ''; headers = None; meal_marker = True
        if meal_marker: continue
        if any(c in ('생산량','총생산량','생산중특이사항','생산책임자','포장방법및단위') for c in compact[:2]):
            headers = None; menu = ''; continue
        if '메뉴명' in compact and '원재료' in compact:
            headers = {'menu':compact.index('메뉴명'), 'material':compact.index('원재료')}
            for j, c in enumerate(compact):
                if c.startswith('1인량'): headers['portion'] = j
                if c.startswith('사용량'): headers['used'] = j
                if c == '단위': headers['unit'] = j
            kind = 'production'; menu = ''; recognized += 1; continue
        if '식단' in compact and '식재' in compact:
            headers = {'menu':compact.index('식단'), 'material':compact.index('식재')}
            for j,c in enumerate(compact):
                if c.startswith('1인량'): headers['portion'] = j
                if c.lower() == 'kg': headers['used'] = j
                if c == '비고': headers['note'] = j
            kind = 'recipe'; menu = ''; recognized += 1; continue
        if '품목명' in compact and '전일재고' in compact and '사용량' in compact:
            headers = {k:compact.index(v) for k,v in {'material':'품목명','opening':'전일재고','used':'사용량','closing':'재고량','incoming':'입고','unit':'단위','supplier':'업체명'}.items() if v in compact}
            kind = 'stock'; recognized += 1; continue
        if not headers: continue
        def get(key):
            j = headers.get(key)
            return row[j] if j is not None and j < len(row) else None
        if kind != 'stock' and clean(get('menu')):
            menu = clean(get('menu'))
        material = clean(get('material'))
        if not material or material in ('원재료','식재','품목명','합계','소계','생산량','총생산량'): continue
        # Do not bleed the menu table into subsequent production/signature sections.
        if any(x in material for x in ('생산책임자', '작성자', '확인자')): headers = None; continue
        if kind != 'stock' and (not menu or menu in ('기타','메뉴명','식단')): continue
        raw = {k:json_value(get(k)) for k in headers}
        unit_raw = clean(get('unit'))
        # These units are explicit in the recognized headers, except stock where the cell is required.
        unit = unit_raw.lower() if kind == 'stock' else 'kg'
        if unit == '㎏': unit = 'kg'
        vals = {k:number(get(k)) for k in ('portion','used','opening','incoming','closing')}
        missing = [k for k in (('used','opening','incoming','closing') if kind == 'stock' else ('portion','used')) if vals[k] is None]
        notes = []
        if not work_date and kind in ('production','stock'): notes.append('날짜 미확인')
        if unit not in ('kg',): notes.append('단위 미확인')
        if missing: notes.append('수치 미확인: '+', '.join(missing))
        if any(v is not None and v < 0 for v in vals.values()): notes.append('음수 검토')
        if kind == 'stock' and all(vals[k] is not None for k in ('opening','incoming','used','closing')):
            if abs(vals['opening'] + vals['incoming'] - vals['used'] - vals['closing']) > .001: notes.append('재고 수식 불일치')
        records.append({'kind':kind,'sheet':sheet,'row':ri,'work_date':work_date,'meal':meal or sheet,'menu':menu if kind!='stock' else '', 'material':material,'unit':unit,'supplier':clean(get('supplier')),**vals,'raw':raw,'columns':headers.copy(),'issues':notes})
    return records, preview, recognized

def extract(job):
    fid, filename, group = job
    path = Path(filename)
    data = {'id':fid,'sheets':[],'records':[],'status':'indexed','error':''}
    try:
        count = 0
        kind = {'조리계획서':'recipe','생산일지':'production','원료수불부':'stock'}.get(group)
        # A cooking plan can also be filed in the school-menu folder.
        is_plan = '조리계획' in clean(path.name)
        for title, rows in workbook_sheets(path):
            if kind or is_plan:
                rec, prev, recognized = parse_sheet(rows,kind or 'recipe', title)
                data['records'].extend(rec); count += recognized
            else:
                prev=[]
                for ri,row in rows:
                    if ri>12: break
                    prev.append({'row':ri,'cells':[json_value(v) for v in row[:24]]})
            data['sheets'].append({'name':title,'preview':prev})
        if kind or is_plan:
            data['status'] = 'extracted' if data['records'] else ('empty_template' if count else 'unrecognized')
        else: data['status'] = 'preview_only'
    except Exception as e:
        data['status'] = 'error'; data['error'] = f'{type(e).__name__}: {str(e)[:300]}'
        # Partial values are never counted as a complete file.
        data['records'] = []
    return data

SCHEMA = '''
CREATE TABLE meta(key TEXT PRIMARY KEY,value TEXT NOT NULL);
CREATE TABLE files(id TEXT PRIMARY KEY,path TEXT NOT NULL,name TEXT,category TEXT,ext TEXT,bytes INTEGER,sha256 TEXT,canonical_id TEXT,status TEXT,error TEXT,sheets TEXT);
CREATE TABLE records(id INTEGER PRIMARY KEY,file_id TEXT NOT NULL REFERENCES files(id),kind TEXT,sheet TEXT,row INTEGER,work_date TEXT,meal TEXT,menu TEXT,material TEXT,unit TEXT,supplier TEXT,portion REAL,used REAL,opening REAL,incoming REAL,closing REAL,raw TEXT,columns TEXT,issues TEXT);
CREATE INDEX record_kind ON records(kind,work_date);
CREATE INDEX record_material ON records(material);
CREATE INDEX record_menu ON records(menu);
CREATE INDEX record_file ON records(file_id);
'''

def build(source, workers=4):
    source = source.resolve(); DB.parent.mkdir(parents=True,exist_ok=True)
    tmp = DB.with_suffix('.tmp'); tmp.unlink(missing_ok=True)
    db = sqlite3.connect(tmp); db.executescript(SCHEMA)
    db.execute('PRAGMA foreign_keys=ON')
    seen, jobs, manifests = {}, [], []
    paths = sorted((p for p in source.rglob('*') if p.is_file() and p.name != '.DS_Store'), key=lambda p:str(p.relative_to(source)))
    # Keep canonical choice deterministic; templates are never preferred to actual data.
    paths.sort(key=lambda p:('복사' in clean(p.name),p.name.startswith('~$'),str(p.relative_to(source))))
    for p in paths:
        rel = str(p.relative_to(source)); fid = hashlib.sha256(rel.encode()).hexdigest()[:20]
        h = hashlib.sha256()
        with p.open('rb') as f:
            for block in iter(lambda:f.read(1024*1024),b''): h.update(block)
        digest=h.hexdigest(); group=category(p); ext=p.suffix.lower()
        status = 'indexed'; canonical = fid
        if p.name.startswith('~$'): status='temporary'
        elif digest in seen: canonical=seen[digest]; status='duplicate'
        elif '복사해서' in clean(p.name): status='template'; seen[digest]=fid
        else:
            seen[digest]=fid
            if ext in ('.xlsx','.xlsb'): jobs.append((fid,str(p),group)); status='pending'
            elif ext in ('.hwp','.xls'): status='conversion_required'
            elif ext == '.pdf': status='ocr_not_run'
            else: status='indexed'
        db.execute('INSERT INTO files VALUES (?,?,?,?,?,?,?,?,?,?,?)',(fid,rel,clean(p.name),group,ext,p.stat().st_size,digest,canonical,status,'','[]'))
        manifests.append((rel,digest))
    db.commit()
    print(f'Inventory: {len(paths)} files; parsing {len(jobs)} representative Excel workbooks',flush=True)
    def save(result):
        db.execute('UPDATE files SET status=?,error=?,sheets=? WHERE id=?',(result['status'],result['error'],json.dumps(result['sheets'],ensure_ascii=False),result['id']))
        for rec in result['records']:
            keys=['kind','sheet','row','work_date','meal','menu','material','unit','supplier','portion','used','opening','incoming','closing','raw','columns','issues']
            vals=[json.dumps(rec[k],ensure_ascii=False) if k in ('raw','columns','issues') else rec[k] for k in keys]
            db.execute('INSERT INTO records(file_id,'+','.join('"'+k+'"' for k in keys)+') VALUES ('+','.join('?' for _ in range(18))+')',[result['id']]+vals)
    with ProcessPoolExecutor(max_workers=workers) as pool:
        for i,result in enumerate(pool.map(extract,jobs,chunksize=1),1):
            save(result)
            if i%40 == 0 or i==len(jobs):
                db.commit();print(f'Parsed {i}/{len(jobs)} · records {db.execute("SELECT count(*) FROM records").fetchone()[0]:,}',flush=True)
    dataset_id=hashlib.sha256(json.dumps(sorted(manifests),ensure_ascii=False).encode()).hexdigest()
    meta={'source_root':str(source),'created_at':datetime.now(timezone.utc).isoformat(),'dataset_id':dataset_id,'parser_version':PARSER_VERSION,'basis':'analysis_snapshot','file_count':len(paths)}
    for k,v in meta.items(): db.execute('INSERT INTO meta VALUES (?,?)',(k,json.dumps(v,ensure_ascii=False)))
    db.commit()
    integrity=db.execute('PRAGMA integrity_check').fetchone()[0]
    if integrity!='ok': raise RuntimeError(integrity)
    counts=dict(db.execute('SELECT kind,count(*) FROM records GROUP BY kind'))
    states=dict(db.execute('SELECT status,count(*) FROM files GROUP BY status'))
    db.close();tmp.replace(DB)
    report={**meta,'records':counts,'file_states':states,'integrity':integrity}
    (ROOT/'reports/ingestion.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    print(json.dumps(report,ensure_ascii=False,indent=2),flush=True)
    return report

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--source',type=Path);p.add_argument('--workers',type=int,default=4);args=p.parse_args()
    build(args.source or source_root(),args.workers)
